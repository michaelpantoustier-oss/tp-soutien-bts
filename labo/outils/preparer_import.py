"""Prépare les fichiers CSV à importer dans le classeur privé du suivi labo.

Les données de l'école ne sont jamais versées dans ce dépôt public : ce script
lit les classeurs Excel en local et écrit les CSV dans un dossier de sortie.

    python3 preparer_import.py --biomed BIOMED_SPG.xlsx --bioalc BIOALC_SPG.xlsx \
        --inventaire INVENTAIRE.xlsx --commandes COMMANDES_LABO.xlsx --sortie ./import

Fichiers produits (une feuille du classeur chacun, même nom) :
  Programme.csv    séquences des scénarios pédagogiques avec TP et matériel repéré
  Articles.csv     inventaire de départ
  Fournisseurs.csv fournisseurs déjà utilisés, avec les catégories achetées
"""
import argparse
import csv
import os
import re
import unicodedata
from datetime import datetime, date

import openpyxl

# Matériel repéré dans le texte des activités : (motif, article, catégorie).
MATERIEL = [
    (r"spectrophotom", "Spectrophotomètre", "Équipement"),
    (r"\bcuves?\b", "Cuves de spectrophotométrie", "Consommable"),
    (r"biuret", "Réactif de biuret", "Réactif / coffret"),
    (r"god[- ]?p[oa]", "Coffret glucose GOD-POD/PAP", "Réactif / coffret"),
    (r"glucose|glycémie", "Coffret glucose", "Réactif / coffret"),
    (r"\burée\b", "Coffret urée", "Réactif / coffret"),
    (r"créatinine", "Coffret créatinine", "Réactif / coffret"),
    (r"cholest", "Coffret cholestérol", "Réactif / coffret"),
    (r"triglyc", "Coffret triglycérides", "Réactif / coffret"),
    (r"acide urique", "Coffret acide urique", "Réactif / coffret"),
    (r"alat|asat|transaminase", "Coffrets ALAT / ASAT", "Réactif / coffret"),
    (r"\bpal\b|phosphatase alcaline", "Coffret phosphatases alcalines", "Réactif / coffret"),
    (r"\bldh\b", "Coffret LDH", "Réactif / coffret"),
    (r"bilirubine", "Coffret bilirubine", "Réactif / coffret"),
    (r"\bciq\b|contrôle de qualité interne|sérum de contrôle|exatrol", "Sérums de contrôle (CIQ) deux niveaux", "Réactif / coffret"),
    (r"électrophorèse|electrophorese", "Kit d'électrophorèse (gel, tampon, colorant)", "Réactif / coffret"),
    (r"chromatograph|clhp|ccm", "Matériel de chromatographie", "Petit matériel"),
    (r"titrage|burette", "Burettes et solutions titrantes", "Verrerie"),
    (r"\bph[- ]?mètre|\bph\b", "pH-mètre et tampons d'étalonnage", "Équipement"),
    (r"conductim", "Conductimètre", "Équipement"),
    (r"balance|pesée", "Balance analytique", "Équipement"),
    (r"verrerie jaugée|fiole|pipette jaugée", "Verrerie jaugée", "Verrerie"),
    (r"micropipette|pipetage", "Micropipettes et cônes", "Petit matériel"),
    (r"centrifug", "Centrifugeuse et tubes", "Équipement"),
    (r"microscop", "Microscopes", "Équipement"),
    (r"\blames?\b|lamelle", "Lames et lamelles", "Consommable"),
    (r"huile à immersion|immersion", "Huile à immersion", "Consommable"),
    (r"\bgram\b", "Kit coloration de Gram", "Réactif / coffret"),
    (r"ziehl|auramine", "Kit Ziehl-Neelsen / auramine", "Réactif / coffret"),
    (r"\bmgg\b|may[- ]grünwald|giemsa", "Colorants MGG / Giemsa", "Réactif / coffret"),
    (r"frottis", "Lames et colorants pour frottis", "Consommable"),
    (r"malassez|thoma|kova|hématimètre|numération", "Cellules de numération (Malassez, Thoma, Kova)", "Petit matériel"),
    (r"bleu trypan", "Bleu trypan", "Réactif / coffret"),
    (r"hémostase|\btq\b|\btca\b|fibrinog", "Réactifs d'hémostase (TQ, TCA, fibrinogène)", "Réactif / coffret"),
    (r"groupage|beth[- ]vincent|simonin|phénotyp|rhésus|agglutination", "Réactifs d'immuno-hématologie (sérums tests, hématies tests)", "Réactif / coffret"),
    (r"elisa|immunodosage", "Kit ELISA et microplaques", "Réactif / coffret"),
    (r"tpha|trod", "Tests immunologiques rapides (TPHA, TROD)", "Réactif / coffret"),
    (r"\bpcr\b|qpcr|thermocycleur", "Réactifs PCR et consommables (thermocycleur)", "Réactif / coffret"),
    (r"extraction.{0,20}(adn|acides nucléiques)|\badn\b", "Kit d'extraction d'ADN", "Réactif / coffret"),
    (r"agarose", "Agarose et tampon d'électrophorèse ADN", "Réactif / coffret"),
    (r"ensemencement|isolement|repiquage", "Anses / öses, écouvillons, becs Bunsen", "Petit matériel"),
    (r"gélose|milieux? de culture|\bmilieux?\b", "Milieux de culture (géloses, bouillons)", "Milieu de culture"),
    (r"drigalski|\bemb\b|chapman|mueller|hektoen|\bsabouraud|columbia|chocolat|\bpca\b|vrbg|\btbx\b|baird|slanetz|chromog|can2", "Géloses spécifiques citées dans la séquence", "Milieu de culture"),
    (r"boîtes? de pétri|boites? de petri", "Boîtes de Petri stériles", "Consommable"),
    (r"\bapi\b|galerie", "Galeries d'identification biochimique", "Réactif / coffret"),
    (r"oxydase", "Bandelettes oxydase", "Réactif / coffret"),
    (r"catalase|peroxyde", "Eau oxygénée (catalase)", "Réactif / coffret"),
    (r"antibiogramme|disques? d'antibiotiques", "Disques d'antibiotiques et Mueller-Hinton", "Réactif / coffret"),
    (r"hémoculture", "Flacons d'hémoculture", "Consommable"),
    (r"souches?", "Souches bactériennes de référence", "Souche / échantillon pédagogique"),
    (r"parasit|helminthe|protozo|kystes?|œufs|paludisme|plasmodium", "Échantillons parasitologiques pédagogiques", "Souche / échantillon pédagogique"),
    (r"levures?|candida|fongique|mycolog", "Souches et milieux de mycologie", "Souche / échantillon pédagogique"),
    (r"filtration sur membrane|rampe de filtration|membranes? filtrantes", "Rampe de filtration et membranes", "Équipement"),
    (r"dénombrement|dilutions? décimales", "Tubes et diluant pour dilutions décimales", "Consommable"),
    (r"étuve|incubat", "Étuves", "Équipement"),
    (r"autoclave|stérilis", "Autoclave", "Équipement"),
    (r"\bpsm\b|poste de sécurité microbiologique|hotte", "Poste de sécurité microbiologique", "Équipement"),
    (r"bain[- ]marie", "Bain-marie", "Équipement"),
    (r"culture cellulaire|lignée|trypsine|dmem|rpmi", "Milieux et consommables de culture cellulaire", "Milieu de culture"),
    (r"spermatozo|sperme|paillettes|gamètes", "Paillettes de sperme bovin et matériel de spermiologie", "Souche / échantillon pédagogique"),
    (r"urinaire|bandelettes? urinaires|ecbu", "Bandelettes urinaires et matériel d'ECBU", "Consommable"),
    (r"ponction veineuse|prélèvement sanguin|bras de simulation", "Matériel de prélèvement (bras de simulation, aiguilles, tubes)", "Consommable"),
    (r"\bepi\b|gants|blouse|lunettes", "EPI (gants, lunettes, blouses)", "EPI / hygiène"),
    (r"dasri|déchets", "Collecteurs DASRI", "EPI / hygiène"),
    (r"désinfect|antisep", "Désinfectants de surface", "EPI / hygiène"),
    (r"\blait\b", "Échantillons de lait (matrice d'étude)", "Souche / échantillon pédagogique"),
    (r"\beau\b.{0,20}(process|potable|analyse)", "Échantillons d'eau (matrice d'étude)", "Souche / échantillon pédagogique"),
    (r"automate|semi-automate", "Automate / semi-automate (temps machine, réactifs dédiés)", "Équipement"),
    (r"maldi", "Accès MALDI-TOF (sortie ou démonstration)", "Équipement"),
]

# Mots indiquant une activité pratique : sans eux, la séquence est un cours ou un TD.
PRATIQUE = re.compile(r"\btp\b|paillasse|manipul|dosage|ensemencement|coloration|frottis|mesure|culture|"
                      r"isolement|extraction|électrophorèse|numération|démonstration|atelier|lecture|pesée|"
                      r"pipetage|centrifug|titrage|dénombrement|groupage|prélèvement", re.I)

# Matières sans manipulation au laboratoire de TP.
HORS_LABO = re.compile(r"anglais|culture g|cge|math|stats|relations professionnelles|comm-projet|projet", re.I)


def norm(s):
    return "" if s is None else re.sub(r"\s+", " ", str(s)).strip()


def dates(texte):
    trouvees = re.findall(r"(\d{2})/(\d{2})/(\d{4})", texte or "")
    iso = ["%s-%s-%s" % (a, m, j) for j, m, a in trouvees]
    return (iso[0], iso[-1]) if iso else ("", "")


def materiel(texte):
    t = texte.lower()
    vus = []
    for motif, nom, _ in MATERIEL:
        if re.search(motif, t) and nom not in vus:
            vus.append(nom)
    return vus


def programme_biomed(chemin):
    wb = openpyxl.load_workbook(chemin, data_only=True)
    lignes = []
    for ws in wb.worksheets:
        if not ws.title.startswith("SPG") or HORS_LABO.search(ws.title):
            continue
        annee = "1A" if "1A" in ws.title else "2A"
        matiere = ws.title.split("·", 1)[-1].strip()
        n = 0
        for r in ws.iter_rows(min_row=5, values_only=True):
            if not r[9]:
                continue
            tc, renf, ress = norm(r[14]), norm(r[15]), norm(r[16])
            if not PRATIQUE.search(tc + " " + renf):
                continue
            mat = materiel(tc + " " + renf + " " + ress + " " + norm(r[9]))
            if not mat:
                continue
            n += 1
            debut, fin = dates(norm(r[7]))
            lignes.append({
                "id": "BM-%s-%s-%02d" % (annee, re.sub(r"\W+", "", matiere)[:12], n),
                "filiere": "BIOMED", "annee": annee, "matiere": matiere, "bloc": norm(r[1]).split("·")[0].strip(),
                "semaines": norm(r[4]), "debut": debut, "fin": fin, "titre": norm(r[9]),
                "tp": tc, "renforcement": renf, "materiel": " | ".join(mat),
            })
    return lignes


def programme_bioalc(chemin):
    wb = openpyxl.load_workbook(chemin, data_only=True)
    lignes = []
    for titre in ("6. SPG 1A détaillé", "7. SPG 2A détaillé"):
        if titre not in wb.sheetnames:
            continue
        annee = "1A" if "1A" in titre else "2A"
        n = 0
        for r in wb[titre].iter_rows(min_row=5, values_only=True):
            if not r[7] or not r[1] or HORS_LABO.search(norm(r[0])):
                continue
            tp_h = r[16] if isinstance(r[16], (int, float)) else 0
            tc, renf = norm(r[12]), norm(r[17])
            if not tp_h and not PRATIQUE.search(tc + " " + renf):
                continue
            mat = materiel(tc + " " + renf + " " + norm(r[7]))
            if not mat:
                continue
            n += 1
            debut, fin = dates(norm(r[6]))
            lignes.append({
                "id": "AC-%s-%02d" % (annee, n), "filiere": "BIOALC", "annee": annee, "matiere": norm(r[0]),
                "bloc": norm(r[1]), "semaines": norm(r[5]), "debut": debut, "fin": fin, "titre": norm(r[7]),
                "tp": tc, "renforcement": re.sub(r"^Mer\.[^\n]*?—\s*[\d h]+\s*\([^)]*\)\s*", "", norm(r[17])),
                "materiel": " | ".join(mat),
            })
    return lignes


def articles(chemin):
    wb = openpyxl.load_workbook(chemin, data_only=True)
    ws = wb.worksheets[0]
    sortie, col = [], None
    for r in ws.iter_rows(values_only=True):
        cells = [norm(c) for c in r]
        if "REACTIF" in cells:
            col = {k: cells.index(k) for k in ("REACTIF", "REFERENCE", "QUANTITE", "EXPIRATION") if k in cells}
            continue
        if not col or not cells[col["REACTIF"]]:
            continue
        ref = r[col["REFERENCE"]]
        ref = str(int(ref)) if isinstance(ref, float) and ref.is_integer() else norm(ref)
        exp = r[col["EXPIRATION"]]
        peremption = exp.strftime("%Y-%m-%d") if isinstance(exp, (datetime, date)) else ""
        notes = "" if peremption or norm(exp) in ("", "/") else "Conservation : " + norm(exp)
        q = r[col["QUANTITE"]] if isinstance(r[col["QUANTITE"]], (int, float)) else 0
        nom = cells[col["REACTIF"]]
        sortie.append({
            "id": "A-IMPORT-%03d" % (len(sortie) + 1), "filiere": "Commun", "categorie": "Réactif / coffret",
            "designation": nom, "reference": ref,
            "fournisseur": "", "quantite": int(q) if float(q).is_integer() else q, "unite": "coffret", "seuil": 1,
            "peremption": peremption, "emplacement": "", "notes": notes,
            "maj_le": datetime.now().strftime("%Y-%m-%d"), "maj_par": "Import inventaire Excel", "actif": "oui",
        })
    return sortie


def sans_accents(s):
    return "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")


def fournisseurs(chemin):
    wb = openpyxl.load_workbook(chemin, data_only=True)
    ws = wb.worksheets[0]
    vus = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        nom, designation = norm(r[2]), norm(r[0])
        if not nom or sans_accents(nom).upper().startswith("FOURNI"):
            continue
        cle = sans_accents(nom).lower().replace("sordabal", "sordalab")
        f = vus.setdefault(cle, {"nom": nom.strip().title() if nom.islower() or nom.isupper() else nom.strip(), "cats": set(), "articles": []})
        for motif, _, cat in MATERIEL:
            if re.search(motif, designation.lower()):
                f["cats"].add(cat)
        if len(f["articles"]) < 6:
            f["articles"].append(designation[:60])
    return [{"nom": f["nom"], "domaine": "Laboratoire", "categories": " | ".join(sorted(f["cats"])) or "Petit matériel",
             "contact": "", "email": "", "telephone": "",
             "notes": "Déjà commandé : " + " ; ".join(f["articles"])} for f in vus.values()]


def ecrire(dossier, nom, colonnes, lignes):
    os.makedirs(dossier, exist_ok=True)
    chemin = os.path.join(dossier, nom)
    with open(chemin, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=colonnes, extrasaction="ignore")
        w.writeheader()
        for l in lignes:
            w.writerow({k: ("'" + v if isinstance(v, str) and v and v[0] in "=+-@" else v) for k, v in l.items()})
    print("%s : %d lignes" % (chemin, len(lignes)))


def main():
    p = argparse.ArgumentParser()
    p.add_argument("--biomed")
    p.add_argument("--bioalc")
    p.add_argument("--inventaire")
    p.add_argument("--commandes")
    p.add_argument("--sortie", default="import")
    a = p.parse_args()
    prog = (programme_biomed(a.biomed) if a.biomed else []) + (programme_bioalc(a.bioalc) if a.bioalc else [])
    ecrire(a.sortie, "Programme.csv", ["id", "filiere", "annee", "matiere", "bloc", "semaines", "debut", "fin", "titre", "tp", "renforcement", "materiel"], prog)
    if a.inventaire:
        ecrire(a.sortie, "Articles.csv", ["id", "filiere", "categorie", "designation", "reference", "fournisseur", "quantite", "unite", "seuil", "peremption", "emplacement", "notes", "maj_le", "maj_par", "actif"], articles(a.inventaire))
    if a.commandes:
        ecrire(a.sortie, "Fournisseurs.csv", ["nom", "domaine", "categories", "contact", "email", "telephone", "notes"], fournisseurs(a.commandes))


if __name__ == "__main__":
    main()
